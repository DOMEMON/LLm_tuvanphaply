"""Generic, bounded table profiling and deterministic query execution.

This experimental core has no domain names, dataset IDs, benchmark answers or
column-name rules. Original cells and physical locators remain available.
"""
from collections import Counter
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import math
from pathlib import Path
import re
import threading
import zipfile

import duckdb
from openpyxl import load_workbook
import sqlglot
from sqlglot import exp


def sha(data):
    return hashlib.sha256(data).hexdigest()


def identifier(name):
    return '"' + name.replace('"', '""') + '"'


def infer(values):
    if not values:
        return 'VARCHAR', lambda x: x
    if all(re.fullmatch(r'-?(0|[1-9]\d*)', v) for v in values):
        if all(-(2**63) <= int(v) < 2**63 for v in values):
            return 'BIGINT', int
        return 'VARCHAR', lambda x: x
    if all(re.fullmatch(r'-?(?:(?:0|[1-9]\d*)(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?', v) for v in values):
        if all(math.isfinite(float(v)) for v in values):
            return 'DOUBLE', float
    if all(re.fullmatch(r'\d{4}-\d{2}-\d{2}', v) for v in values):
        try:
            for v in values:
                date.fromisoformat(v)
            return 'DATE', lambda x: x
        except ValueError:
            pass
    return 'VARCHAR', lambda x: x


class Table:
    def __init__(self, path, *, missing_values=('',), sheet=None):
        self.path = Path(path)
        if self.path.stat().st_size > 20 * 1024 * 1024:
            raise ValueError('FILE_TOO_LARGE')
        with self.path.open('rb') as stream:
            raw = stream.read(20 * 1024 * 1024 + 1)
        if len(raw) > 20 * 1024 * 1024:
            raise ValueError('FILE_TOO_LARGE')
        self.hash = sha(raw)
        if self.path.suffix.lower() == '.csv':
            text = raw.decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(text[:65536], delimiters=',;\t|')
            except csv.Error:
                if any(separator in text[:65536] for separator in ',;\t|'):
                    raise ValueError('UNRECOGNIZED_CSV_DIALECT')
                dialect = csv.excel  # One-column CSV has no separator to infer.
            reader = csv.reader(io.StringIO(text, newline=''), dialect)
            records = []
            locators = []
            previous_line = 0
            for row in reader:
                if row:
                    if len(records) >= 100001 or len(row) > 256 or any(len(v) > 16000 for v in row):
                        raise ValueError('TABLE_LIMIT')
                    records.append(row)
                    locators.append({'line_start': previous_line + 1, 'line_end': reader.line_num})
                previous_line = reader.line_num
            self.format = {'type': 'csv', 'delimiter': dialect.delimiter}
        elif self.path.suffix.lower() == '.xlsx':
            try:
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    if len(archive.infolist()) > 1024 or sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
                        raise ValueError('WORKBOOK_EXPANSION_LIMIT')
            except zipfile.BadZipFile as exc:
                raise ValueError('INVALID_XLSX_ARCHIVE') from exc
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                if sheet is None and len(book.sheetnames) != 1:
                    raise ValueError('MULTIPLE_SHEETS_REQUIRE_SELECTION')
                if sheet and sheet not in book.sheetnames:
                    raise ValueError('UNKNOWN_WORKSHEET')
                ws = book[sheet or book.sheetnames[0]]
                if (ws.max_column and ws.max_column > 256) or (ws.max_row and ws.max_row > 100001):
                    raise ValueError('TABLE_LIMIT')
                records, locators = [], []
                for index, row in enumerate(ws.iter_rows(), 1):
                    if index > 100001 or len(row) > 256:
                        raise ValueError('TABLE_LIMIT')
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('FORMULAS_REQUIRE_PRECOMPUTED_VALUES')
                    values = [v.value for v in row]
                    if not any(v is not None for v in values):
                        continue
                    if any(len(str(v)) > 16000 for v in values if v is not None):
                        raise ValueError('TABLE_LIMIT')
                    records.append(['' if v is None else v.isoformat() if isinstance(v, (date, datetime))
                                    else str(v) for v in values])
                    locators.append({'sheet': ws.title, 'row': index})
                self.format = {'type': 'xlsx', 'sheet': ws.title}
            finally:
                book.close()
        else:
            raise ValueError('SUPPORTED_FORMATS_CSV_XLSX')
        if len(records) < 2:
            raise ValueError('EMPTY_TABLE')
        # Spreadsheet exports may carry formatted but wholly empty columns at
        # the right edge. Ignore only columns whose header AND every cell are
        # empty; a blank header above real data remains an invalid table.
        while (records[0] and not records[0][-1].strip() and
               all(len(row) == len(records[0]) and not row[-1].strip() for row in records[1:])):
            records = [row[:-1] for row in records]
        self.headers, self.raw_rows = records[0], records[1:]
        self.locators = locators[1:]
        if (len(self.headers) > 256 or len(self.raw_rows) > 100000
                or any(not h.strip() for h in self.headers)
                or len({h.casefold() for h in self.headers}) != len(self.headers)
                or '_row_id' in {h.casefold() for h in self.headers}):
            raise ValueError('INVALID_HEADERS_OR_TABLE_LIMIT')
        if any(len(row) != len(self.headers) or any(len(v) > 16000 for v in row) for row in self.raw_rows):
            raise ValueError('RAGGED_OR_OVERSIZED_ROW')
        self.rows = [{'_row_id': i + 1} for i in range(len(self.raw_rows))]
        self.columns = []
        for index, header in enumerate(self.headers):
            values = [row[index] for row in self.raw_rows]
            present = [v for v in values if v not in missing_values]
            kind, cast = infer(present)
            counts = Counter(present)
            column = {'name': header, 'type': kind, 'null_count': len(values) - len(present),
                      'distinct_count': len(counts), 'unique_non_null': len(present) == len(values) == len(counts),
                      'examples': [v for v, _ in sorted(counts.items(), key=lambda p: (-p[1], p[0]))[:6]]}
            if kind in {'BIGINT', 'DOUBLE', 'DATE'} and present:
                column.update(min=min(map(cast, present)), max=max(map(cast, present)))
            self.columns.append(column)
            for record, value in zip(self.rows, values):
                record[header] = None if value in missing_values else cast(value)
        self.types = {'_row_id': 'BIGINT', **{c['name']: c['type'] for c in self.columns}}
        self.profile = {'row_count': len(self.rows), 'columns': self.columns,
                        'missing_values': list(missing_values), 'format': self.format,
                        'source_sha256': self.hash,
                        'row_identity': '_row_id is the 1-based source record index; not a business entity ID.'}
        self.db = duckdb.connect(':memory:', config={'enable_external_access': False,
                        'autoload_known_extensions': False, 'autoinstall_known_extensions': False,
                        'threads': 1, 'memory_limit': '256MB'})
        declarations = ', '.join(f'{identifier(k)} {v}' for k, v in self.types.items())
        self.db.execute('CREATE TABLE data (' + declarations + ')')
        self.db.execute('BEGIN TRANSACTION')
        try:
            self.db.executemany('INSERT INTO data VALUES (' + ','.join('?' for _ in self.types) + ')',
                                [list(row.values()) for row in self.rows])
            self.db.execute('COMMIT')
        except Exception:
            self.db.execute('ROLLBACK')
            self.db.close()
            raise
        self.db.execute('SET lock_configuration=true')

    def close(self):
        self.db.close()

    def sql(self, sql, params=None):
        timer = threading.Timer(3, self.db.interrupt)
        timer.daemon = True
        timer.start()
        try:
            cur = self.db.execute(sql, params or [])
            names = [c[0] for c in cur.description]
            rows = cur.fetchmany(1001)
            if len(rows) > 1000:
                raise ValueError('RESULT_LIMIT')
            return [dict(zip(names, [v.isoformat() if isinstance(v, (date, datetime)) else v for v in row])) for row in rows]
        finally:
            timer.cancel()

    def provenance(self, row_ids, columns):
        return [{'row_id': i, 'columns': columns, 'locator': self.locators[i - 1],
                 'source_sha256': self.hash} for i in row_ids if 1 <= i <= len(self.rows)]


SQL_NODES = {'select', 'column', 'identifier', 'table', 'from', 'where', 'and', 'or', 'not',
             'eq', 'neq', 'gt', 'gte', 'lt', 'lte', 'literal', 'null', 'is', 'in', 'between',
             'like', 'ilike', 'paren', 'alias', 'order', 'ordered', 'limit', 'group', 'distinct',
             'count', 'sum', 'avg', 'min', 'max', 'star', 'neg', 'boolean', 'lower', 'upper'}


def validate_sql(sql, table):
    if len(sql) > 8000:
        raise ValueError('SQL_LENGTH')
    statements = sqlglot.parse(sql, read='duckdb')
    if len(statements) != 1 or not isinstance(statements[0], exp.Select):
        raise ValueError('SINGLE_SELECT_REQUIRED')
    tree = statements[0]
    for node in tree.walk():
        if node.key not in SQL_NODES:
            raise ValueError('UNSUPPORTED_SQL_NODE:' + node.key)
        if isinstance(node, exp.Table) and (node.name != 'data' or node.db or node.catalog or node.alias):
            raise ValueError('UNKNOWN_TABLE')
        if isinstance(node, exp.Column) and (node.name not in table.types or node.table not in ('', 'data')):
            raise ValueError('UNKNOWN_COLUMN')
    if len(list(tree.find_all(exp.Table))) != 1:
        raise ValueError('ONE_TABLE_REQUIRED')
    return tree.sql(dialect='duckdb')


OPS = {'eq': '=', 'ne': '<>', 'gt': '>', 'gte': '>=', 'lt': '<', 'lte': '<='}


def compile_plan(task, table):
    """Compile checked typed predicates; neither model SQL nor Python is executed."""
    if task['status'] != 'answer':
        return None
    operation = task['operation']
    field = task['field']
    if operation not in ('count', 'sum', 'avg', 'min', 'max', 'select', 'distinct'):
        raise ValueError('INVALID_OPERATION')
    if field and field not in table.types:
        raise ValueError('UNKNOWN_COLUMN')
    if operation in ('sum', 'avg') and table.types.get(field) not in ('BIGINT', 'DOUBLE'):
        raise ValueError('NUMERIC_AGGREGATE_REQUIRES_NUMERIC_FIELD')
    selected = task['columns']
    if any(col not in table.types for col in selected):
        raise ValueError('UNKNOWN_COLUMN')
    if task['filter_logic'] not in ('and', 'or'):
        raise ValueError('FILTER_LOGIC')
    clauses, params, used = [], [], set(selected)
    if field:
        used.add(field)
    for predicate in task['filters']:
        column, operator, value = predicate['column'], predicate['operator'], predicate['value']
        if column not in table.types:
            raise ValueError('UNKNOWN_COLUMN')
        used.add(column)
        name = identifier(column)
        if operator in ('is_null', 'not_null'):
            clauses.append(name + (' IS NULL' if operator == 'is_null' else ' IS NOT NULL'))
            continue
        kind = table.types[column]
        if operator == 'contains':
            if kind != 'VARCHAR':
                raise ValueError('CONTAINS_REQUIRES_TEXT')
            clauses.append(f'contains(lower({name}), lower(?))')
        elif operator in OPS:
            if kind == 'BIGINT':
                try:
                    number = Decimal(value)
                except InvalidOperation as exc:
                    raise ValueError('INTEGER_VALUE_REQUIRED') from exc
                if not number.is_finite() or number != number.to_integral_value() or not -(2**63) <= number < 2**63:
                    raise ValueError('INTEGER_VALUE_REQUIRED')
                value = int(number)
            elif kind == 'DOUBLE':
                value = float(value)
                if not math.isfinite(value):
                    raise ValueError('FINITE_VALUE_REQUIRED')
            elif kind == 'DATE':
                date.fromisoformat(value)
            clauses.append(name + ' ' + OPS[operator] + ' ?')
        else:
            raise ValueError('UNKNOWN_OPERATOR')
        params.append(value)
    where = ' WHERE ' + (' AND ' if task['filter_logic'] == 'and' else ' OR ').join(clauses) if clauses else ''
    if operation in ('select', 'distinct'):
        if not selected:
            raise ValueError('SELECT_REQUIRES_COLUMNS')
        projection = ('DISTINCT ' if operation == 'distinct' else '') + ', '.join(map(identifier, selected))
    else:
        if operation != 'count' and not field:
            raise ValueError('AGGREGATE_FIELD_REQUIRED')
        projection = operation.upper() + '(' + (identifier(field) if field else '*') + ') AS result'
    sql = 'SELECT ' + projection + ' FROM data' + where
    order_by = task['order_by']
    if order_by:
        if operation not in ('select', 'distinct') or order_by not in table.types:
            raise ValueError('INVALID_ORDER')
        direction = task['direction']
        if direction not in ('asc', 'desc'):
            raise ValueError('INVALID_DIRECTION')
        sql += ' ORDER BY ' + identifier(order_by) + ' ' + direction + ' NULLS LAST'
    limit = task['limit']
    if operation in ('select', 'distinct'):
        if not 1 <= limit <= 1000:
            raise ValueError('INVALID_LIMIT')
        sql += ' LIMIT ' + str(limit)
    evidence_sql = 'SELECT _row_id FROM data' + where
    return sql, params, evidence_sql, sorted(used)


def execute_plan(task, table):
    compiled = compile_plan(task, table)
    if compiled is None:
        return {'status': task['status'], 'data': [], 'evidence': []}
    sql, params, evidence_sql, used = compiled
    data = table.sql(sql, params)
    # Aggregate lineage is a replayable full-set predicate, not top-k citations.
    evidence = {'source_sha256': table.hash, 'columns': used, 'query': sql, 'params': params,
                'selection_query': evidence_sql,
                'matching_rows': table.sql('SELECT COUNT(*) AS result FROM (' + evidence_sql + ')', params)[0]['result']}
    return {'status': 'answer', 'data': data, 'evidence': evidence}
