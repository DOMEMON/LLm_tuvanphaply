import json

from openpyxl import Workbook
import pytest

from .table import Table, execute_plan, infer, validate_sql


def plan(**kwargs):
    return {'status': 'answer', 'operation': 'count', 'field': '', 'columns': [],
            'filter_logic': 'and', 'filters': [], 'order_by': '', 'direction': 'asc', 'limit': 1000, **kwargs}


@pytest.fixture
def table(tmp_path):
    path = tmp_path / 'data.csv'
    path.write_text('code,amount,category,day\n001,10,A,2026-01-01\n002,,B,2026-01-02\n003,30,A,2026-01-03\n', encoding='utf-8')
    t = Table(path)
    yield t
    t.close()


def test_preserve_identifier_and_null(table):
    assert table.types['code'] == 'VARCHAR'
    assert table.rows[0]['code'] == '001'
    assert table.rows[1]['amount'] is None
    assert table.types['day'] == 'DATE'
    assert execute_plan(plan(operation='avg', field='amount'), table)['data'] == [{'result': 20}]


def test_trim_only_wholly_empty_trailing_export_columns(tmp_path):
    from openpyxl import Workbook

    file = tmp_path / 'export.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.append(['name', 'fee', None, None])
    sheet.append(['Permit A', 0, None, None])
    sheet.append(['Permit B', 8, None, None])
    book.save(file)
    source = Table(file)
    try:
        assert source.headers == ['name', 'fee']
        assert source.profile['row_count'] == 2
        assert source.provenance([1], ['name'])[0]['locator']['row'] == 2
    finally:
        source.close()

    sheet.cell(row=3, column=3, value='unlabelled information')
    book.save(file)
    with pytest.raises(ValueError, match='INVALID_HEADERS'):
        Table(file)


def test_count_rows_and_count_non_null_differ(table):
    assert execute_plan(plan(), table)['data'] == [{'result': 3}]
    assert execute_plan(plan(field='amount'), table)['data'] == [{'result': 2}]


def test_predicate_lineage(table):
    result = execute_plan(plan(filters=[{'column': 'amount', 'operator': 'gte', 'value': '20'}]), table)
    assert result['data'] == [{'result': 1}]
    assert result['evidence']['matching_rows'] == 1
    assert result['evidence']['source_sha256'] == table.hash
    assert table.provenance([3], ['amount'])[0]['locator']['line_start'] == 4


def test_no_sql_injection_from_plan_value(table):
    result = execute_plan(plan(filters=[{'column': 'category', 'operator': 'eq', 'value': "A' OR 1=1 --"}]), table)
    assert result['data'] == [{'result': 0}]


@pytest.mark.parametrize('sql', ["SELECT * FROM read_csv('/etc/passwd')", 'DROP TABLE data',
    'SELECT * FROM data; SELECT * FROM data', 'SELECT * FROM other',
    'SELECT * FROM data JOIN data d ON true', "SELECT getenv('HOME') FROM data",
    'SELECT * FROM (SELECT * FROM data)', 'SELECT missing FROM data'])
def test_reject_unbounded_or_external_sql(table, sql):
    with pytest.raises(ValueError):
        validate_sql(sql, table)


def test_sql_supported_aggregate(table):
    sql = validate_sql('SELECT COUNT(*) AS result FROM data WHERE amount IS NOT NULL', table)
    assert table.sql(sql) == [{'result': 2}]


def test_dont_sum_text(table):
    with pytest.raises(ValueError):
        execute_plan(plan(operation='sum', field='code'), table)


def test_xlsx_matches_csv(table, tmp_path):
    book = Workbook()
    ws = book.active
    ws.append(table.headers)
    for row in table.raw_rows:
        ws.append(row)
    path = tmp_path / 'test.xlsx'
    book.save(path)
    x = Table(path)
    try:
        assert x.rows == table.rows
        assert x.provenance([1], ['code'])[0]['locator'] == {'sheet': 'Sheet', 'row': 2}
    finally:
        x.close()


def test_xlsx_formula_rejected(tmp_path):
    book = Workbook()
    book.active.append(['name', 'value'])
    book.active.append(['x', '=1+1'])
    path = tmp_path / 'formula.xlsx'
    book.save(path)
    with pytest.raises(ValueError, match='FORMULAS'):
        Table(path)


def test_duplicate_headers_rejected(tmp_path):
    path = tmp_path / 'bad.csv'
    path.write_text('a,A\n1,2\n', encoding='utf-8')
    with pytest.raises(ValueError, match='HEADERS'):
        Table(path)


def test_numeric_notation_without_changing_padded_identifiers():
    assert infer(['10.', '20.5', '.5'])[0] == 'DOUBLE'
    assert infer(['001', '002'])[0] == 'VARCHAR'


def test_large_integer_filters_are_exact(tmp_path):
    path = tmp_path / 'large.csv'
    path.write_text('key,value\n9007199254740992,1\n9007199254740993,2\n', encoding='utf-8')
    t = Table(path)
    try:
        result = execute_plan(plan(operation='select', columns=['value'], filters=[
            {'column': 'key', 'operator': 'eq', 'value': '9007199254740993'}]), t)
        assert result['data'] == [{'value': 2}]
    finally:
        t.close()


def test_single_column_csv_is_supported(tmp_path):
    path = tmp_path / 'single.csv'
    path.write_text('code\n001\n002\n', encoding='utf-8')
    t = Table(path)
    try:
        assert t.rows == [{'_row_id': 1, 'code': '001'}, {'_row_id': 2, 'code': '002'}]
    finally:
        t.close()
